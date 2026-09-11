"""Сервис работы с MinIO (S3-совместимое хранилище).

Используется для загрузки и хранения файлов прайс-листов и спецификаций.
Один экземпляр на процесс (Singleton), клиент создаётся лениво.
"""

import threading
from typing import Any, ClassVar, Self
from urllib.parse import quote

from botocore.exceptions import ClientError

from app.core.config import Settings


class MinioService:
    """Асинхронный доступ к MinIO/S3-хранилищу через aioboto3."""

    _instances: ClassVar[dict[str, "MinioService"]] = {}
    _lock: ClassVar[threading.Lock] = threading.Lock()

    def __new__(cls, settings: Settings) -> Self:
        key = settings.minio.host
        with cls._lock:
            if key not in cls._instances:
                instance = super().__new__(cls)
                instance._init(settings)
                cls._instances[key] = instance
        return cls._instances[key]

    def _init(self, settings: Settings) -> None:
        self._settings = settings
        self._bucket = settings.minio.bucket
        self._client: Any | None = None
        self._client_ctx: Any | None = None
        self._client_lock = threading.Lock()

    async def _get_client(self) -> Any:
        """Ленивая инициализация aioboto3-клиента.

        session.client() возвращает async context manager — входим в него один
        раз и держим клиент открытым всё время жизни Singleton.
        """
        if self._client is None:
            with self._client_lock:
                if self._client is None:
                    import aioboto3

                    session = aioboto3.Session(
                        aws_access_key_id=self._settings.minio.root_user,
                        aws_secret_access_key=self._settings.minio.root_password,
                    )
                    self._client_ctx = session.client(
                        "s3",
                        endpoint_url=f"http://{self._settings.minio.endpoint}",
                    )
                    self._client = await self._client_ctx.__aenter__()
        return self._client

    async def create_bucket_if_not_exists(self) -> None:
        """Создать бакет, если не существует."""
        client = await self._get_client()
        try:
            await client.head_bucket(Bucket=self._bucket)
        except ClientError:
            await client.create_bucket(Bucket=self._bucket)

    async def upload_file(self, key: str, file_path: str) -> str:
        """Загрузить локальный файл в хранилище.

        Возвращает публичный URL файла.
        """
        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        await client.upload_file(file_path, self._bucket, encoded_key)
        return f"http://{self._settings.minio.endpoint}/{self._bucket}/{encoded_key}"

    async def upload_bytes(self, key: str, file_bytes: bytes, content_type: str = "application/octet-stream") -> str:
        """Загрузить байты файла в хранилище.

        Возвращает публичный URL файла.
        """
        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        await client.put_object(
            Bucket=self._bucket,
            Key=encoded_key,
            Body=file_bytes,
            ContentType=content_type,
        )
        return f"http://{self._settings.minio.endpoint}/{self._bucket}/{encoded_key}"

    async def upload_fileobj(self, key: str, fileobj: Any, content_type: str = "application/octet-stream") -> str:
        """Потоковая загрузка файла в хранилище (multipart, без чтения всего файла в память).

        fileobj — seekable file-like объект (например, SpooledTemporaryFile из
        FastAPI UploadFile.file). Возвращает публичный URL файла.
        """
        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        await client.upload_fileobj(
            Fileobj=fileobj,
            Bucket=self._bucket,
            Key=encoded_key,
            ExtraArgs={"ContentType": content_type},
        )
        return f"http://{self._settings.minio.endpoint}/{self._bucket}/{encoded_key}"

    async def delete_file(self, key: str) -> None:
        """Удалить файл из хранилища."""
        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        await client.delete_object(Bucket=self._bucket, Key=encoded_key)

    async def file_exists(self, key: str) -> bool:
        """Проверить существование файла."""
        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        try:
            await client.head_object(Bucket=self._bucket, Key=encoded_key)
            return True
        except ClientError:
            return False

    async def download_fileobj(self, key: str) -> Any:
        """Скачать файл из хранилища в файл-объект.

        Возвращает BytesIO с содержимым файла.
        """
        import io

        client = await self._get_client()
        encoded_key = quote(key, safe="/")
        response = await client.get_object(Bucket=self._bucket, Key=encoded_key)
        body = await response["Body"].read()
        return io.BytesIO(body)

    async def download_workbook(self, key: str) -> Any:
        """Скачать Excel-файл из хранилища и открыть как openpyxl Workbook.

        Возвращает openpyxl Workbook.
        """
        from openpyxl import load_workbook

        fileobj = await self.download_fileobj(key)
        return load_workbook(fileobj, read_only=True)
